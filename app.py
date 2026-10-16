import io
import os
import json
from datetime import datetime
import streamlit as st
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference

# ---------------------------------------------------------------------------
# CONFIGURAÇÃO GERAL E ESTILO SÓBRIO / ACADÊMICO
# ---------------------------------------------------------------------------
st.set_page_config(
    page_title="VZN EcoClima — Sistema de Análise Microclimática",
    page_icon="📐",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Estilização CSS institucional adaptativa para Tema Claro e Tema Escuro
st.markdown("""
<style>
    body, [class*="css"] {
        font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
    }
    .titulo-institucional {
        font-size: 1.65rem;
        font-weight: 700;
        letter-spacing: -0.02em;
        margin-bottom: 0.2rem;
        color: var(--text-color) !important;
    }
    .subtitulo-institucional {
        font-size: 0.92rem;
        color: var(--text-color) !important;
        opacity: 0.75;
        margin-bottom: 1.0rem;
        border-bottom: 1px solid rgba(128, 128, 128, 0.25);
        padding-bottom: 0.6rem;
    }
    /* Expander adaptativo para tema claro e escuro */
    div[data-testid="stExpander"] {
        border: 1px solid rgba(128, 128, 128, 0.25) !important;
        border-radius: 6px !important;
        background-color: transparent !important;
    }
    div[data-testid="stExpander"] > details {
        background-color: transparent !important;
    }
    button[kind="primary"], button[kind="secondary"] {
        border-radius: 4px !important;
        font-weight: 500 !important;
    }
    .barra-projeto {
        border: 1px solid rgba(128, 128, 128, 0.28);
        border-radius: 8px;
        padding: 0.75rem 1rem;
        margin-bottom: 1rem;
        background-color: rgba(128, 128, 128, 0.04);
    }
</style>
""", unsafe_allow_html=True)

# Caminho do arquivo local onde os projetos salvos ficam armazenados
ARQUIVO_PROJETOS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "projetos_salvos.json")

# ---------------------------------------------------------------------------
# DADOS PRÉ-DEFINIDOS (PADRÃO DO PROJETO)
# ---------------------------------------------------------------------------
PRESETS = {
    "Dia 14 de Agosto": {
        "titulo_dia": "Dia 14 de agosto",
        "titulo_grafico": "DIA 14 DE AGOSTO - BOSQUE",
        "min_y": 24.0,
        "max_y": 33.0,
        "df": pd.DataFrame({
            'Pontos': ['P1', 'P2', 'P3', 'P4', 'P5', 'P6', 'P7', 'P8', 'P9', 'P10'],
            'Manha': [26.2, 26.5, 26.5, 26.9, 27.2, 27.4, 27.1, 27.2, 27.3, 27.6],
            'Tarde': [28.5, 29.0, 29.5, 29.6, 30.4, 29.9, 30.9, 31.4, 30.3, 30.5],
            'Noite': [25.5, 26.1, 26.2, 26.4, 26.9, 26.7, 26.8, 26.9, 27.0, 27.1]
        })
    },
    "Dia 16 de Agosto": {
        "titulo_dia": "Dia 16 de agosto",
        "titulo_grafico": "DIA 16 DE AGOSTO - BOSQUE",
        "min_y": 24.0,
        "max_y": 33.0,
        "df": pd.DataFrame({
            'Pontos': ['P1', 'P2', 'P3', 'P4', 'P5', 'P6', 'P7', 'P8', 'P9', 'P10'],
            'Manha': [26.0, 25.3, 25.3, 26.6, 26.2, 25.5, 25.4, 26.2, 26.0, 26.3],
            'Tarde': [28.5, 28.6, 28.9, 31.0, 30.8, 29.5, 30.4, 30.6, 30.0, 29.9],
            'Noite': [25.8, 26.0, 26.1, 26.7, 26.8, 27.0, 27.0, 26.9, 26.6, 27.1]
        })
    },
    "Dia 19 de Agosto": {
        "titulo_dia": "Dia 19 de agosto",
        "titulo_grafico": "DIA 19 DE AGOSTO - BOSQUE",
        "min_y": 22.0,
        "max_y": 33.0,
        "df": pd.DataFrame({
            'Pontos': ['P1', 'P2', 'P3', 'P4', 'P5', 'P6', 'P7', 'P8', 'P9', 'P10'],
            'Manha': [24.7, 24.0, 23.4, 23.6, 23.8, 23.6, 23.5, 23.8, 24.2, 24.4],
            'Tarde': [28.2, 28.5, 28.9, 31.5, 31.0, 28.9, 31.0, 31.8, 30.0, 30.6],
            'Noite': [27.1, 25.5, 25.3, 26.7, 26.2, 26.1, 25.8, 27.0, 27.1, 25.6]
        })
    }
}

# Paletas de cores disponíveis
PALETAS = {
    "OriginLab Pro (Preto, Vermelho, Azul)": {
        "Manha": "#222222",
        "Tarde": "#e74c3c",
        "Noite": "#2980b9"
    },
    "Acadêmico Monocromático (Escala de Cinza)": {
        "Manha": "#111111",
        "Tarde": "#666666",
        "Noite": "#aaaaaa"
    },
    "Nature / Science (Tons Frios & Quentes)": {
        "Manha": "#2c3e50",
        "Tarde": "#d35400",
        "Noite": "#16a085"
    },
    "Excel Clássico": {
        "Manha": "#265b79",
        "Tarde": "#e27641",
        "Noite": "#236532"
    }
}

FONTES_DISPONIVEIS = [
    "Times New Roman (Recomendado ABNT)",
    "Arial (Científico Internacional)",
    "Cambria (Serifada Elegante)",
    "Garamond (Clássico Acadêmico)",
    "Calibri (Sans-serif Moderna)"
]

POSICOES_LEGENDA = [
    "Superior Direito (Interno)",
    "Superior Esquerdo (Interno)",
    "Inferior Direito (Interno)",
    "Inferior Esquerdo (Interno)",
    "Centro Superior (Externo)",
    "Lateral Direita (Externo)",
    "Personalizada (Sliders X/Y)"
]

DEFAULT_CONFIG = {
    "rotulo_x": "Pontos",
    "rotulo_y": "Temperatura do ar (°C)",
    "opcao_fonte": "Times New Roman (Recomendado ABNT)",
    "tamanho_fonte": 11.0,
    "opcao_paleta": "OriginLab Pro (Preto, Vermelho, Azul)",
    "cor_manha": "#222222",
    "cor_tarde": "#e74c3c",
    "cor_noite": "#2980b9",
    "espessura_linha": 1.2,
    "com_marcadores": False,
    "exibir_grade": False,
    "com_linhas_distintas": False,
    "show_manha": True,
    "show_tarde": True,
    "show_noite": True,
    "exibir_legenda": True,
    "moldura_legenda": True,
    "orientacao_leg": "Vertical (Empilhados)",
    "posicao_preset": "Superior Direito (Interno)",
    "leg_x": 0.98,
    "leg_y": 0.98,
    "borda_top": True,
    "borda_bottom": True,
    "borda_left": True,
    "borda_right": True,
    "ticks_topo": True,
    "ticks_dir": True,
    "ticks_baixo": True,
    "ticks_esq": True,
    "direcao_ticks": "Para dentro (in)",
    "minor_ticks": True,
    "cor_moldura": "#000000",
    "esp_moldura": 1.15,
}


# ---------------------------------------------------------------------------
# FUNÇÕES DE PERSISTÊNCIA (SALVAR E CARREGAR PROJETOS)
# ---------------------------------------------------------------------------
def ler_banco_projetos():
    """Lê o arquivo JSON local contendo os projetos salvos."""
    if not os.path.exists(ARQUIVO_PROJETOS):
        return {"ultimo_projeto_ativo": None, "projetos": {}}
    try:
        with open(ARQUIVO_PROJETOS, "r", encoding="utf-8") as f:
            dados = json.load(f)
            if not isinstance(dados, dict) or "projetos" not in dados:
                return {"ultimo_projeto_ativo": None, "projetos": {}}
            return dados
    except Exception:
        return {"ultimo_projeto_ativo": None, "projetos": {}}


def gravar_banco_projetos(banco):
    """Salva o dicionário de projetos no arquivo JSON local."""
    with open(ARQUIVO_PROJETOS, "w", encoding="utf-8") as f:
        json.dump(banco, f, ensure_ascii=False, indent=2)


def construir_payload_projeto(nome_projeto, df, titulo_dia, titulo_grafico, min_y, max_y, cfg):
    """Empacota todos os dados e configurações atuais do projeto para serialização JSON."""
    df_clean = df.copy()
    registros = []
    for _, row in df_clean.iterrows():
        registros.append({
            "Pontos": str(row.get("Pontos", "")),
            "Manha": None if pd.isna(row.get("Manha")) else float(row.get("Manha")),
            "Tarde": None if pd.isna(row.get("Tarde")) else float(row.get("Tarde")),
            "Noite": None if pd.isna(row.get("Noite")) else float(row.get("Noite")),
        })

    return {
        "nome_projeto": nome_projeto.strip(),
        "atualizado_em": datetime.now().strftime("%d/%m/%Y às %H:%M:%S"),
        "titulo_dia": titulo_dia,
        "titulo_grafico": titulo_grafico,
        "min_y": float(min_y),
        "max_y": float(max_y),
        "dados": registros,
        "config": dict(cfg),
    }


def aplicar_projeto_na_sessao(proj_dict):
    """Restaura um projeto salvo para dentro do st.session_state."""
    st.session_state.nome_projeto_atual = proj_dict.get("nome_projeto", proj_dict.get("titulo_dia", "Meu Projeto"))
    st.session_state.titulo_dia = proj_dict.get("titulo_dia", "Dia 14 de agosto")
    st.session_state.titulo_grafico = proj_dict.get("titulo_grafico", "DIA 14 DE AGOSTO - BOSQUE")
    st.session_state.min_y = float(proj_dict.get("min_y", 24.0))
    st.session_state.max_y = float(proj_dict.get("max_y", 33.0))
    st.session_state.ultimo_salvamento = proj_dict.get("atualizado_em")

    registros = proj_dict.get("dados", [])
    if registros:
        df_loaded = pd.DataFrame(registros)
        for col in ["Pontos", "Manha", "Tarde", "Noite"]:
            if col not in df_loaded.columns:
                df_loaded[col] = np.nan if col != "Pontos" else ""
        st.session_state.df_atual = df_loaded[["Pontos", "Manha", "Tarde", "Noite"]].copy()
    else:
        st.session_state.df_atual = PRESETS["Dia 14 de Agosto"]["df"].copy()

    cfg_salva = proj_dict.get("config", {})
    nova_cfg = DEFAULT_CONFIG.copy()
    nova_cfg.update(cfg_salva)
    st.session_state.cfg = nova_cfg
    st.session_state.editor_version = st.session_state.get("editor_version", 0) + 1


def sanitizar_dados_tabela(df):
    """
    Corrige automaticamente problemas comuns ao colar do Excel em português:
    1. Vírgula que virou zero/centena ao colar (ex: 27,4 -> 274.0 vira 27.4).
    2. Valores None/vazios na coluna Pontos quando linhas extras são coladas.
    """
    modificado = False
    df_novo = df.copy()

    for idx in range(len(df_novo)):
        val_ponto = df_novo.iloc[idx]['Pontos']
        if pd.isna(val_ponto) or str(val_ponto).strip().lower() in ['none', 'nan', '']:
            df_novo.iloc[idx, df_novo.columns.get_loc('Pontos')] = f"P{idx + 1}"
            modificado = True

    for col in ['Manha', 'Tarde', 'Noite']:
        if col in df_novo.columns:
            for idx in range(len(df_novo)):
                val = df_novo.iloc[idx][col]
                if pd.notnull(val):
                    try:
                        v_float = float(val)
                        if 60.0 < v_float <= 600.0:
                            df_novo.iloc[idx, df_novo.columns.get_loc(col)] = round(v_float / 10.0, 1)
                            modificado = True
                        elif 600.0 < v_float <= 6000.0:
                            df_novo.iloc[idx, df_novo.columns.get_loc(col)] = round(v_float / 100.0, 1)
                            modificado = True
                    except (ValueError, TypeError):
                        pass

    return df_novo, modificado


def gerar_figura_matplotlib(df, titulo_grafico, rotulo_x, rotulo_y,
                            min_y, max_y, series_ativas, cores_periodos, estilos_linhas,
                            fonte_nome, tamanho_fonte, espessura_linha, com_marcadores,
                            bordas_config, ticks_config, legenda_config, exibir_grade):
    """
    Renderiza a figura do Matplotlib com controle milimétrico de tipografia,
    cores, linhas visíveis, bordas do quadro e posicionamento livre da legenda.
    """
    # 1. Configuração de Tipografia
    if "Times" in fonte_nome:
        plt.rcParams['font.serif'] = ['Times New Roman', 'Times', 'DejaVu Serif', 'serif']
        plt.rcParams['font.family'] = 'serif'
    elif "Cambria" in fonte_nome:
        plt.rcParams['font.serif'] = ['Cambria', 'Georgia', 'serif']
        plt.rcParams['font.family'] = 'serif'
    elif "Garamond" in fonte_nome:
        plt.rcParams['font.serif'] = ['Garamond', 'Georgia', 'serif']
        plt.rcParams['font.family'] = 'serif'
    else:  # Arial, Helvetica, Sans-Serif
        plt.rcParams['font.sans-serif'] = ['Arial', 'Helvetica', 'DejaVu Sans', 'sans-serif']
        plt.rcParams['font.family'] = 'sans-serif'

    fig, ax = plt.subplots(figsize=(7.2, 4.4), dpi=200)
    fig.patch.set_facecolor('white')
    ax.set_facecolor('white')

    pontos = df['Pontos'].astype(str).tolist()
    x_indices = np.arange(len(pontos))

    marcadores_map = {'Manha': 'o', 'Tarde': 's', 'Noite': '^'}

    # 2. Plotagem das linhas (apenas as séries ativas marcadas pela usuária)
    for periodo in ['Manha', 'Tarde', 'Noite']:
        if periodo in df.columns and periodo in series_ativas:
            m = marcadores_map[periodo] if com_marcadores else None
            ls = estilos_linhas.get(periodo, '-')
            cor = cores_periodos.get(periodo, '#000000')

            ax.plot(
                x_indices, df[periodo],
                label=periodo,
                color=cor,
                linestyle=ls,
                linewidth=espessura_linha,
                marker=m,
                markersize=5,
                markeredgecolor=cor,
                markerfacecolor='white' if com_marcadores and "Monocromático" in str(cores_periodos) else cor
            )

    # 3. Escalas e Limites dos Eixos
    ax.set_xticks(x_indices)
    ax.set_xticklabels(pontos, fontsize=tamanho_fonte - 0.5)

    todas_temps = []
    for col in ['Manha', 'Tarde', 'Noite']:
        if col in df.columns and col in series_ativas:
            todas_temps.extend(pd.to_numeric(df[col], errors='coerce').dropna().tolist())

    if todas_temps:
        real_min = min(todas_temps)
        real_max = max(todas_temps)
        pad_min = (int(real_min // 2) * 2) if min_y is None else min_y
        pad_max = (int(np.ceil(real_max / 2)) * 2) if max_y is None else max_y
        if pad_max <= pad_min:
            pad_max = pad_min + 2
        ax.set_ylim(pad_min, pad_max)

    if len(x_indices) > 0:
        ax.set_xlim(-0.55, len(x_indices) - 0.45)

    # 4. Controle das Bordas (Spines)
    cor_borda = bordas_config.get('cor', 'black')
    esp_borda = bordas_config.get('espessura', 1.15)

    ax.spines['top'].set_visible(bordas_config.get('topo', True))
    ax.spines['right'].set_visible(bordas_config.get('direita', True))
    ax.spines['bottom'].set_visible(bordas_config.get('baixo', True))
    ax.spines['left'].set_visible(bordas_config.get('esquerda', True))

    for s_name in ['top', 'right', 'bottom', 'left']:
        ax.spines[s_name].set_color(cor_borda)
        ax.spines[s_name].set_linewidth(esp_borda)

    # 5. Controle dos Ticks (Traços de Escala)
    direcao_t = ticks_config.get('direcao', 'in')
    ticks_topo = ticks_config.get('topo', True) and bordas_config.get('topo', True)
    ticks_dir = ticks_config.get('direita', True) and bordas_config.get('direita', True)
    ticks_baixo = ticks_config.get('baixo', True) and bordas_config.get('baixo', True)
    ticks_esq = ticks_config.get('esquerda', True) and bordas_config.get('esquerda', True)

    ax.tick_params(
        which='major', direction=direcao_t,
        top=ticks_topo, right=ticks_dir, bottom=ticks_baixo, left=ticks_esq,
        length=5.5, width=1.1, color=cor_borda, labelsize=tamanho_fonte - 0.5
    )

    if ticks_config.get('minor_ticks', True):
        ax.xaxis.set_minor_locator(ticker.AutoMinorLocator(2))
        ax.yaxis.set_major_locator(ticker.MultipleLocator(1))
        ax.yaxis.set_minor_locator(ticker.AutoMinorLocator(2))
        ax.tick_params(
            which='minor', direction=direcao_t,
            top=ticks_topo, right=ticks_dir, bottom=ticks_baixo, left=ticks_esq,
            length=3.2, width=0.8, color=cor_borda
        )
    else:
        ax.yaxis.set_major_locator(ticker.MultipleLocator(2))

    ax.grid(exibir_grade, color='#e0e0e0', linestyle='--', linewidth=0.6)

    # 6. Rótulos e Título
    ax.set_xlabel(rotulo_x, fontsize=tamanho_fonte, color='black', labelpad=8)
    ax.set_ylabel(rotulo_y, fontsize=tamanho_fonte, color='black', labelpad=8)

    if titulo_grafico and titulo_grafico.strip():
        pad_val = 12 if bordas_config.get('topo', True) else 22
        ax.set_title(titulo_grafico, fontsize=tamanho_fonte + 1.5, color='black', pad=pad_val, fontweight='normal')

    # 7. Posicionamento e Customização Livre da Legenda (Arrastável / Coordenadas)
    if legenda_config.get('exibir', True) and len(series_ativas) > 0:
        pos_x = legenda_config.get('pos_x', 0.98)
        pos_y = legenda_config.get('pos_y', 0.98)
        loc_ancora = legenda_config.get('loc', 'upper right')
        ncol_leg = len(series_ativas) if legenda_config.get('orientacao') == 'Horizontal' else 1
        moldura_leg = legenda_config.get('moldura', True)

        leg = ax.legend(
            loc=loc_ancora,
            bbox_to_anchor=(pos_x, pos_y),
            ncol=ncol_leg,
            frameon=moldura_leg,
            edgecolor='black' if moldura_leg else 'none',
            facecolor='white',
            framealpha=1.0,
            fontsize=tamanho_fonte - 1.5,
            handlelength=1.8,
            handletextpad=0.5,
            borderpad=0.55,
            columnspacing=1.5
        )
        if moldura_leg:
            leg.get_frame().set_linewidth(1.0)

    plt.tight_layout()
    return fig


def gerar_excel_bytes(df, titulo_dia, titulo_grafico, min_y, max_y):
    """Gera a planilha Excel (.xlsx) com a tabela formatada e o gráfico nativo ao lado."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dados e Grafico"

    borda_fina = Border(
        left=Side(style='thin', color='A6A6A6'),
        right=Side(style='thin', color='A6A6A6'),
        top=Side(style='thin', color='A6A6A6'),
        bottom=Side(style='thin', color='A6A6A6')
    )

    ws.merge_cells("B2:E2")
    header = ws.cell(row=2, column=2, value=titulo_dia)
    header.font = Font(name="Times New Roman", size=11, bold=True)
    header.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(2, 6):
        ws.cell(row=2, column=c).border = borda_fina

    colunas = ["Pontos", "Manha", "Tarde", "Noite"]
    for idx, col in enumerate(colunas, start=2):
        c = ws.cell(row=3, column=idx, value=col)
        c.font = Font(name="Times New Roman", size=10, bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = borda_fina

    for r_idx, row in df.iterrows():
        r = 4 + r_idx
        ws.cell(row=r, column=2, value=str(row['Pontos'])).alignment = Alignment(horizontal="center")
        for col_idx, col_name in enumerate(["Manha", "Tarde", "Noite"], start=3):
            val = row.get(col_name, np.nan)
            cell = ws.cell(row=r, column=col_idx, value=float(val) if pd.notnull(val) else None)
            cell.alignment = Alignment(horizontal="center")
            cell.border = borda_fina
            if cell.value is not None:
                cell.number_format = '0.0'
        ws.cell(row=r, column=2).border = borda_fina

    fim_dados = 3 + len(df)

    chart = LineChart()
    chart.title = titulo_grafico
    chart.style = 10
    chart.y_axis.title = "Temperatura do ar (°C)"
    chart.x_axis.title = "Pontos"
    if min_y is not None:
        chart.y_axis.scaling.min = min_y
    if max_y is not None:
        chart.y_axis.scaling.max = max_y
    chart.y_axis.majorUnit = 2
    chart.width = 16
    chart.height = 9.5

    dados_ref = Reference(ws, min_col=3, min_row=3, max_col=5, max_row=fim_dados)
    cats_ref = Reference(ws, min_col=2, min_row=4, max_row=fim_dados)
    chart.add_data(dados_ref, titles_from_data=True)
    chart.set_categories(cats_ref)
    ws.add_chart(chart, "G2")

    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 11
    ws.column_dimensions['C'].width = 11
    ws.column_dimensions['D'].width = 11
    ws.column_dimensions['E'].width = 11
    ws.column_dimensions['F'].width = 4

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# INICIALIZAÇÃO DO ESTADO DA SESSÃO (CARREGA ÚLTIMO PROJETO SALVO SE EXISTIR)
# ---------------------------------------------------------------------------
banco_projetos = ler_banco_projetos()

if "editor_version" not in st.session_state:
    st.session_state.editor_version = 0

if "inicializado" not in st.session_state:
    st.session_state.inicializado = True
    ultimo_ativo = banco_projetos.get("ultimo_projeto_ativo")
    if ultimo_ativo and ultimo_ativo in banco_projetos.get("projetos", {}):
        aplicar_projeto_na_sessao(banco_projetos["projetos"][ultimo_ativo])
        st.session_state.ultimo_preset = "Selecione..."
    else:
        st.session_state.ultimo_preset = "Dia 14 de Agosto"
        st.session_state.nome_projeto_atual = "Dia 14 de agosto"
        st.session_state.titulo_dia = PRESETS["Dia 14 de Agosto"]["titulo_dia"]
        st.session_state.titulo_grafico = PRESETS["Dia 14 de Agosto"]["titulo_grafico"]
        st.session_state.min_y = PRESETS["Dia 14 de Agosto"]["min_y"]
        st.session_state.max_y = PRESETS["Dia 14 de Agosto"]["max_y"]
        st.session_state.df_atual = PRESETS["Dia 14 de Agosto"]["df"].copy()
        st.session_state.cfg = DEFAULT_CONFIG.copy()
        st.session_state.ultimo_salvamento = None

v = st.session_state.editor_version
cfg = st.session_state.cfg

# ---------------------------------------------------------------------------
# BARRA LATERAL: CONFIGURAÇÕES E ESTILO
# ---------------------------------------------------------------------------
with st.sidebar:
    st.markdown("### **VZN EcoClima**")
    st.caption("Sistema de Análise Microclimática")
    st.markdown("---")

    st.markdown("#### **Tipografia Acadêmica**")
    idx_fonte = FONTES_DISPONIVEIS.index(cfg["opcao_fonte"]) if cfg["opcao_fonte"] in FONTES_DISPONIVEIS else 0
    opcao_fonte = st.selectbox(
        "Família de Fonte:",
        FONTES_DISPONIVEIS,
        index=idx_fonte,
        key=f"fonte_{v}"
    )
    tamanho_fonte = st.slider(
        "Tamanho da Fonte:",
        min_value=9.0, max_value=14.0,
        value=float(cfg["tamanho_fonte"]),
        step=0.5,
        key=f"tam_fonte_{v}"
    )

    st.markdown("---")
    st.markdown("#### **Cores das Séries**")
    opcoes_paleta_lista = list(PALETAS.keys()) + ["Personalizada (Definir manualmente)"]
    idx_paleta = opcoes_paleta_lista.index(cfg["opcao_paleta"]) if cfg["opcao_paleta"] in opcoes_paleta_lista else 0
    opcao_paleta = st.selectbox(
        "Paleta de Cores:",
        opcoes_paleta_lista,
        index=idx_paleta,
        key=f"paleta_{v}"
    )

    if opcao_paleta == "Personalizada (Definir manualmente)":
        c_p1, c_p2, c_p3 = st.columns(3)
        with c_p1:
            cor_manha = st.color_picker("Manhã", value=cfg["cor_manha"], key=f"cor_m_{v}")
        with c_p2:
            cor_tarde = st.color_picker("Tarde", value=cfg["cor_tarde"], key=f"cor_t_{v}")
        with c_p3:
            cor_noite = st.color_picker("Noite", value=cfg["cor_noite"], key=f"cor_n_{v}")
        cores_selecionadas = {"Manha": cor_manha, "Tarde": cor_tarde, "Noite": cor_noite}
    else:
        cor_manha = cfg["cor_manha"]
        cor_tarde = cfg["cor_tarde"]
        cor_noite = cfg["cor_noite"]
        cores_selecionadas = PALETAS[opcao_paleta]

    st.markdown("---")
    st.markdown("#### **Traços e Marcadores**")
    espessura_linha = st.slider(
        "Espessura do Traço (pt):",
        min_value=0.8, max_value=2.8,
        value=float(cfg["espessura_linha"]),
        step=0.1,
        key=f"esp_linha_{v}"
    )

    col_opt1, col_opt2 = st.columns(2)
    with col_opt1:
        com_marcadores = st.checkbox("Marcadores nos pontos", value=bool(cfg["com_marcadores"]), key=f"marc_{v}")
    with col_opt2:
        exibir_grade = st.checkbox("Linhas de grade", value=bool(cfg["exibir_grade"]), key=f"grade_{v}")

    com_linhas_distintas = st.checkbox(
        "Linhas distintas (Sólido / Tracejado / Pontilhado)",
        value=bool(cfg["com_linhas_distintas"]),
        help="Ideal para impressões acadêmicas em preto e branco.",
        key=f"linhas_dist_{v}"
    )
    if com_linhas_distintas:
        estilos_linhas = {"Manha": "-", "Tarde": "--", "Noite": ":"}
    else:
        estilos_linhas = {"Manha": "-", "Tarde": "-", "Noite": "-"}

    st.markdown("---")
    st.markdown("#### **Origem dos Dados / Modelos**")
    opcoes_preset = ["Selecione...", "Criar Novo / Em Branco"] + list(PRESETS.keys())
    preset_escolhido = st.selectbox(
        "Carregar modelo pré-cadastrado:",
        opcoes_preset,
        index=0,
        key=f"preset_sel_{v}"
    )

    uploaded_file = st.file_uploader(
        "Importar planilha (.xlsx ou .csv):",
        type=["xlsx", "csv"],
        key=f"up_excel_{v}"
    )

    uploaded_proj_file = st.file_uploader(
        "Importar arquivo de projeto (.json):",
        type=["json"],
        help="Carregue um arquivo .json de projeto salvo anteriormente no computador.",
        key=f"up_json_{v}"
    )

# Tratamento de carregamento de modelo pré-cadastrado
if preset_escolhido != "Selecione...":
    if preset_escolhido in PRESETS:
        p = PRESETS[preset_escolhido]
        st.session_state.nome_projeto_atual = p["titulo_dia"]
        st.session_state.titulo_dia = p["titulo_dia"]
        st.session_state.titulo_grafico = p["titulo_grafico"]
        st.session_state.min_y = p["min_y"]
        st.session_state.max_y = p["max_y"]
        st.session_state.df_atual = p["df"].copy()
        st.session_state.ultimo_salvamento = None
    elif preset_escolhido == "Criar Novo / Em Branco":
        st.session_state.nome_projeto_atual = "Novo Projeto"
        st.session_state.titulo_dia = "Novo Dia"
        st.session_state.titulo_grafico = "NOVO DIA - BOSQUE"
        st.session_state.min_y = 20.0
        st.session_state.max_y = 35.0
        st.session_state.df_atual = pd.DataFrame({
            'Pontos': [f"P{i}" for i in range(1, 11)],
            'Manha': [25.0] * 10,
            'Tarde': [30.0] * 10,
            'Noite': [26.0] * 10
        })
        st.session_state.ultimo_salvamento = None
    st.session_state.editor_version += 1
    st.rerun()

# Tratamento de importação de arquivo de projeto (.json)
if uploaded_proj_file is not None:
    try:
        proj_importado = json.load(uploaded_proj_file)
        if isinstance(proj_importado, dict) and "dados" in proj_importado:
            aplicar_projeto_na_sessao(proj_importado)
            # Também salva no banco local para ficar disponível na lista de projetos
            nome_imp = st.session_state.nome_projeto_atual
            banco_projetos["projetos"][nome_imp] = proj_importado
            banco_projetos["ultimo_projeto_ativo"] = nome_imp
            gravar_banco_projetos(banco_projetos)
            st.toast(f"Projeto '{nome_imp}' importado e salvo com sucesso!", icon="✅")
            st.rerun()
        else:
            st.sidebar.error("Arquivo JSON não possui o formato válido de projeto EcoClima.")
    except Exception as e:
        st.sidebar.error(f"Erro ao importar projeto: {e}")

# Tratamento de importação de planilha (.xlsx ou .csv)
if uploaded_file is not None:
    try:
        if uploaded_file.name.endswith(".csv"):
            df_up = pd.read_csv(uploaded_file)
        else:
            df_up = pd.read_excel(uploaded_file)

        col_map = {}
        for c in df_up.columns:
            c_clean = str(c).strip().lower()
            if 'ponto' in c_clean:
                col_map[c] = 'Pontos'
            elif 'manh' in c_clean or 'manha' in c_clean:
                col_map[c] = 'Manha'
            elif 'tarde' in c_clean:
                col_map[c] = 'Tarde'
            elif 'noite' in c_clean:
                col_map[c] = 'Noite'

        df_up = df_up.rename(columns=col_map)
        colunas_necessarias = ['Pontos', 'Manha', 'Tarde', 'Noite']
        if all(c in df_up.columns for c in colunas_necessarias):
            st.session_state.df_atual = df_up[colunas_necessarias].copy()
            nome_base = os.path.splitext(uploaded_file.name)[0]
            st.session_state.nome_projeto_atual = nome_base
            st.session_state.titulo_dia = nome_base
            st.session_state.titulo_grafico = f"{nome_base.upper()} - BOSQUE"
            st.session_state.editor_version += 1
            st.toast(f"Planilha '{uploaded_file.name}' carregada com sucesso.", icon="✅")
            st.rerun()
        else:
            st.sidebar.error("A planilha precisa ter as colunas: Pontos, Manha, Tarde e Noite.")
    except Exception as e:
        st.sidebar.error(f"Erro ao ler planilha: {e}")

# ---------------------------------------------------------------------------
# CABEÇALHO INSTITUCIONAL SÓBRIO
# ---------------------------------------------------------------------------
st.markdown('<div class="titulo-institucional">VZN EcoClima — Sistema de Análise Microclimática</div>', unsafe_allow_html=True)
st.markdown('<div class="subtitulo-institucional">Laboratório de Monitoramento Térmico e Produção de Gráficos Científicos | Projeto Acadêmico</div>', unsafe_allow_html=True)

# Contêiner no topo para a Barra de Salvamento e Gerenciamento de Projetos
container_barra_projeto = st.container()

# ---------------------------------------------------------------------------
# ÁREA DE TRABALHO
# ---------------------------------------------------------------------------
col_tabela, col_grafico = st.columns([1, 1.25], gap="large")

with col_tabela:
    st.markdown("#### Matriz de Dados Experimentais")

    # Ferramenta para colar do Excel
    with st.expander("Importar Bloco de Dados do Excel (Copiar e Colar)", expanded=False):
        st.caption("Copie as células no Excel (Ctrl+C) e cole no campo abaixo. As vírgulas são convertidas automaticamente.")
        destino_colar = st.selectbox(
            "Destino dos dados:",
            ["Apenas Manhã", "Apenas Tarde", "Apenas Noite", "Manhã, Tarde e Noite (3 colunas)", "Tabela Completa (Pontos + 3 colunas)"]
        )
        texto_colado = st.text_area("Área de transferência:", height=85, placeholder="Cole aqui os dados copiados do Excel...")

        if st.button("Aplicar Dados na Tabela", use_container_width=True):
            if texto_colado.strip():
                linhas_raw = [l.strip() for l in texto_colado.strip().split('\n') if l.strip()]
                df_temp = st.session_state.df_atual.copy()

                if destino_colar in ["Apenas Manhã", "Apenas Tarde", "Apenas Noite"]:
                    col_alvo = "Manha" if "Manhã" in destino_colar else ("Tarde" if "Tarde" in destino_colar else "Noite")
                    valores = []
                    for l in linhas_raw:
                        val_str = l.replace(',', '.').strip()
                        try:
                            v = float(val_str)
                            if 60.0 < v <= 600.0:
                                v = v / 10.0
                            valores.append(v)
                        except ValueError:
                            valores.append(np.nan)

                    if len(valores) > len(df_temp):
                        linhas_extras = len(valores) - len(df_temp)
                        novos_pontos = [f"P{len(df_temp) + i + 1}" for i in range(linhas_extras)]
                        df_extra = pd.DataFrame({'Pontos': novos_pontos, 'Manha': [np.nan]*linhas_extras, 'Tarde': [np.nan]*linhas_extras, 'Noite': [np.nan]*linhas_extras})
                        df_temp = pd.concat([df_temp, df_extra], ignore_index=True)

                    for i, v in enumerate(valores):
                        df_temp.at[i, col_alvo] = v

                elif destino_colar == "Manhã, Tarde e Noite (3 colunas)":
                    novas_linhas = []
                    for l in linhas_raw:
                        partes = [p.strip().replace(',', '.') for p in l.replace(';', '\t').split('\t')]
                        nums = []
                        for p in partes[:3]:
                            try:
                                v = float(p)
                                if 60.0 < v <= 600.0:
                                    v = v / 10.0
                                nums.append(v)
                            except ValueError:
                                nums.append(np.nan)
                        while len(nums) < 3:
                            nums.append(np.nan)
                        novas_linhas.append(nums)

                    df_temp = pd.DataFrame(novas_linhas, columns=['Manha', 'Tarde', 'Noite'])
                    df_temp.insert(0, 'Pontos', [f"P{i+1}" for i in range(len(df_temp))])

                elif destino_colar == "Tabela Completa (Pontos + 3 colunas)":
                    novas_linhas = []
                    for l in linhas_raw:
                        partes = [p.strip().replace(',', '.') for p in l.replace(';', '\t').split('\t')]
                        ponto_nome = partes[0]
                        nums = []
                        for p in partes[1:4]:
                            try:
                                v = float(p)
                                if 60.0 < v <= 600.0:
                                    v = v / 10.0
                                nums.append(v)
                            except ValueError:
                                nums.append(np.nan)
                        while len(nums) < 3:
                            nums.append(np.nan)
                        novas_linhas.append([ponto_nome] + nums)

                    df_temp = pd.DataFrame(novas_linhas, columns=['Pontos', 'Manha', 'Tarde', 'Noite'])

                st.session_state.df_atual = df_temp
                st.session_state.editor_version += 1
                st.success("Dados colados e formatados com sucesso.")
                st.rerun()

    # Identificadores e Eixos
    t_col1, t_col2 = st.columns(2)
    with t_col1:
        titulo_dia_input = st.text_input(
            "Identificador da Coleta / Dia:",
            value=st.session_state.titulo_dia,
            key=f"tit_dia_{v}"
        )
    with t_col2:
        titulo_grafico_input = st.text_input(
            "Título do Gráfico:",
            value=st.session_state.titulo_grafico,
            key=f"tit_graf_{v}"
        )

    r_col1, r_col2 = st.columns(2)
    with r_col1:
        rotulo_x_input = st.text_input(
            "Legenda Eixo X:",
            value=cfg["rotulo_x"],
            key=f"rot_x_{v}"
        )
    with r_col2:
        rotulo_y_input = st.text_input(
            "Legenda Eixo Y:",
            value=cfg["rotulo_y"],
            key=f"rot_y_{v}"
        )

    # Limites da escala
    e_col1, e_col2, e_col3, e_col4 = st.columns([1, 1, 1, 1])
    with e_col1:
        min_y_val = st.number_input(
            "Y Mínimo (°C):",
            value=float(st.session_state.min_y),
            step=1.0,
            key=f"min_y_{v}"
        )
    with e_col2:
        max_y_val = st.number_input(
            "Y Máximo (°C):",
            value=float(st.session_state.max_y),
            step=1.0,
            key=f"max_y_{v}"
        )

    # Grade de edição interativa
    editor_key = f"editor_principal_{v}"
    df_bruto = st.data_editor(
        st.session_state.df_atual,
        num_rows="dynamic",
        use_container_width=True,
        column_config={
            "Pontos": st.column_config.TextColumn("Pontos", width="small", required=True),
            "Manha": st.column_config.NumberColumn("Manha (°C)", format="%.1f", step=0.1),
            "Tarde": st.column_config.NumberColumn("Tarde (°C)", format="%.1f", step=0.1),
            "Noite": st.column_config.NumberColumn("Noite (°C)", format="%.1f", step=0.1),
        },
        height=370,
        key=editor_key
    )

    df_sanitizado, foi_corrigido = sanitizar_dados_tabela(df_bruto)
    if foi_corrigido:
        st.session_state.df_atual = df_sanitizado
        st.session_state.editor_version += 1
        st.toast("Formatação decimal ajustada automaticamente.", icon="ℹ️")
        st.rerun()
    else:
        df_editado = df_sanitizado
        st.session_state.df_atual = df_editado.copy()

    with e_col3:
        st.write("")
        st.write("")
        if st.button("Renumerar P1..Pn", use_container_width=True):
            st.session_state.df_atual['Pontos'] = [f"P{i+1}" for i in range(len(st.session_state.df_atual))]
            st.session_state.editor_version += 1
            st.rerun()
    with e_col4:
        st.write("")
        st.write("")
        if st.button("Padronizar Decimais", use_container_width=True):
            df_corr, _ = sanitizar_dados_tabela(st.session_state.df_atual)
            st.session_state.df_atual = df_corr
            st.session_state.editor_version += 1
            st.rerun()

with col_grafico:
    # -----------------------------------------------------------------------
    # CONTROLE DE SÉRIES ATIVAS (QUAIS LINHAS APARECEM NO GRÁFICO)
    # -----------------------------------------------------------------------
    st.markdown("#### Visualização Gráfica Interativa")

    col_ser_label, col_ser_m, col_ser_t, col_ser_n = st.columns([1.5, 1, 1, 1])
    with col_ser_label:
        st.markdown("**Séries visíveis:**")
    with col_ser_m:
        show_manha = st.checkbox("Manhã", value=bool(cfg["show_manha"]), key=f"sh_m_{v}")
    with col_ser_t:
        show_tarde = st.checkbox("Tarde", value=bool(cfg["show_tarde"]), key=f"sh_t_{v}")
    with col_ser_n:
        show_noite = st.checkbox("Noite", value=bool(cfg["show_noite"]), key=f"sh_n_{v}")

    series_ativas = []
    if show_manha:
        series_ativas.append("Manha")
    if show_tarde:
        series_ativas.append("Tarde")
    if show_noite:
        series_ativas.append("Noite")

    # -----------------------------------------------------------------------
    # CONFIGURAÇÕES AVANÇADAS: LEGENDA (POSIÇÃO LIVRE / ARRASTAR) E BORDAS
    # -----------------------------------------------------------------------
    with st.expander("📍 **Posicionamento da Legenda (Arrastar / Ajustar Posição e Formato)**", expanded=False):
        c_leg1, c_leg2 = st.columns(2)
        with c_leg1:
            exibir_legenda = st.checkbox("Exibir legenda", value=bool(cfg["exibir_legenda"]), key=f"ex_leg_{v}")
            moldura_legenda = st.checkbox("Caixa de contorno na legenda (Padrão OriginLab)", value=bool(cfg["moldura_legenda"]), key=f"mold_leg_{v}")
            opcoes_orient = ["Vertical (Empilhados)", "Horizontal (Lado a lado)"]
            idx_orient = opcoes_orient.index(cfg["orientacao_leg"]) if cfg["orientacao_leg"] in opcoes_orient else 0
            orientacao_leg = st.radio("Orientação dos itens:", opcoes_orient, index=idx_orient, key=f"ori_leg_{v}")
        with c_leg2:
            idx_pos_preset = POSICOES_LEGENDA.index(cfg["posicao_preset"]) if cfg["posicao_preset"] in POSICOES_LEGENDA else 0
            posicao_preset = st.selectbox(
                "Posição Rápida:",
                POSICOES_LEGENDA,
                index=idx_pos_preset,
                key=f"pos_pre_{v}"
            )

        # Mapeamento ou ajuste milimétrico dos sliders
        if posicao_preset == "Superior Direito (Interno)":
            pad_x, pad_y, loc_ancora = 0.98, 0.98, 'upper right'
        elif posicao_preset == "Superior Esquerdo (Interno)":
            pad_x, pad_y, loc_ancora = 0.02, 0.98, 'upper left'
        elif posicao_preset == "Inferior Direito (Interno)":
            pad_x, pad_y, loc_ancora = 0.98, 0.05, 'lower right'
        elif posicao_preset == "Inferior Esquerdo (Interno)":
            pad_x, pad_y, loc_ancora = 0.02, 0.05, 'lower left'
        elif posicao_preset == "Centro Superior (Externo)":
            pad_x, pad_y, loc_ancora = 0.50, 1.15, 'upper center'
        elif posicao_preset == "Lateral Direita (Externo)":
            pad_x, pad_y, loc_ancora = 1.02, 0.50, 'center left'
        else:
            pad_x, pad_y, loc_ancora = 0.98, 0.98, 'upper right'

        # Se o preset atual é o mesmo salvo na configuração, preserva a coordenada fina ajustada
        val_x_inicial = float(cfg["leg_x"]) if posicao_preset == cfg["posicao_preset"] else float(pad_x)
        val_y_inicial = float(cfg["leg_y"]) if posicao_preset == cfg["posicao_preset"] else float(pad_y)

        st.markdown("**Ajuste Milimétrico de Posição (Arraste os controles para mover a legenda):**")
        slider_x_col, slider_y_col = st.columns(2)
        with slider_x_col:
            leg_x = st.slider("Posição Horizontal (X):", min_value=-0.1, max_value=1.2, value=val_x_inicial, step=0.01, key=f"leg_x_{v}_{posicao_preset}")
        with slider_y_col:
            leg_y = st.slider("Posição Vertical (Y):", min_value=-0.1, max_value=1.3, value=val_y_inicial, step=0.01, key=f"leg_y_{v}_{posicao_preset}")

        legenda_config = {
            'exibir': exibir_legenda,
            'pos_x': leg_x,
            'pos_y': leg_y,
            'loc': loc_ancora,
            'orientacao': orientacao_leg.split(' ')[0],
            'moldura': moldura_legenda
        }

    with st.expander("📐 **Bordas e Ticks do Quadro (Escolher quais linhas aparecem)**", expanded=False):
        st.markdown("**Quais bordas da moldura devem aparecer:**")
        b_c1, b_c2, b_c3, b_c4 = st.columns(4)
        with b_c1:
            borda_top = st.checkbox("Borda Superior", value=bool(cfg["borda_top"]), key=f"b_top_{v}")
        with b_c2:
            borda_bottom = st.checkbox("Borda Inferior", value=bool(cfg["borda_bottom"]), key=f"b_bot_{v}")
        with b_c3:
            borda_left = st.checkbox("Borda Esquerda", value=bool(cfg["borda_left"]), key=f"b_left_{v}")
        with b_c4:
            borda_right = st.checkbox("Borda Direita", value=bool(cfg["borda_right"]), key=f"b_right_{v}")

        st.markdown("**Quais traços de escala (ticks) devem aparecer:**")
        t_c1, t_c2, t_c3 = st.columns(3)
        with t_c1:
            ticks_topo = st.checkbox("Ticks no Topo", value=bool(cfg["ticks_topo"]), key=f"t_top_{v}")
            ticks_dir = st.checkbox("Ticks na Direita", value=bool(cfg["ticks_dir"]), key=f"t_dir_{v}")
        with t_c2:
            ticks_baixo = st.checkbox("Ticks na Base", value=bool(cfg["ticks_baixo"]), key=f"t_bot_{v}")
            ticks_esq = st.checkbox("Ticks na Esquerda", value=bool(cfg["ticks_esq"]), key=f"t_esq_{v}")
        with t_c3:
            opcoes_dir_ticks = ["Para dentro (in)", "Para fora (out)"]
            idx_dir_ticks = opcoes_dir_ticks.index(cfg["direcao_ticks"]) if cfg["direcao_ticks"] in opcoes_dir_ticks else 0
            direcao_ticks = st.radio("Direção dos traços:", opcoes_dir_ticks, index=idx_dir_ticks, key=f"dir_t_{v}")
            minor_ticks = st.checkbox("Traços menores (minor ticks)", value=bool(cfg["minor_ticks"]), key=f"min_t_{v}")

        cor_moldura = st.color_picker("Cor das Bordas:", value=cfg["cor_moldura"], key=f"cor_mold_{v}")
        esp_moldura = st.slider("Espessura das Bordas (pt):", min_value=0.5, max_value=2.5, value=float(cfg["esp_moldura"]), step=0.05, key=f"esp_mold_{v}")

        bordas_config = {
            'topo': borda_top,
            'baixo': borda_bottom,
            'esquerda': borda_left,
            'direita': borda_right,
            'cor': cor_moldura,
            'espessura': esp_moldura
        }

        ticks_config = {
            'topo': ticks_topo,
            'direita': ticks_dir,
            'baixo': ticks_baixo,
            'esquerda': ticks_esq,
            'direcao': 'in' if 'dentro' in direcao_ticks else 'out',
            'minor_ticks': minor_ticks
        }

    # Sincroniza o estado atual das configurações na sessão
    titulo_dia_anterior = st.session_state.titulo_dia
    if titulo_dia_input != titulo_dia_anterior:
        if st.session_state.get("nome_projeto_atual") in (titulo_dia_anterior, "Novo Projeto", "Novo Dia"):
            st.session_state.nome_projeto_atual = titulo_dia_input

    st.session_state.titulo_dia = titulo_dia_input
    st.session_state.titulo_grafico = titulo_grafico_input
    st.session_state.min_y = min_y_val
    st.session_state.max_y = max_y_val
    st.session_state.cfg = {
        "rotulo_x": rotulo_x_input,
        "rotulo_y": rotulo_y_input,
        "opcao_fonte": opcao_fonte,
        "tamanho_fonte": tamanho_fonte,
        "opcao_paleta": opcao_paleta,
        "cor_manha": cor_manha,
        "cor_tarde": cor_tarde,
        "cor_noite": cor_noite,
        "espessura_linha": espessura_linha,
        "com_marcadores": com_marcadores,
        "exibir_grade": exibir_grade,
        "com_linhas_distintas": com_linhas_distintas,
        "show_manha": show_manha,
        "show_tarde": show_tarde,
        "show_noite": show_noite,
        "exibir_legenda": exibir_legenda,
        "moldura_legenda": moldura_legenda,
        "orientacao_leg": orientacao_leg,
        "posicao_preset": posicao_preset,
        "leg_x": leg_x,
        "leg_y": leg_y,
        "borda_top": borda_top,
        "borda_bottom": borda_bottom,
        "borda_left": borda_left,
        "borda_right": borda_right,
        "ticks_topo": ticks_topo,
        "ticks_dir": ticks_dir,
        "ticks_baixo": ticks_baixo,
        "ticks_esq": ticks_esq,
        "direcao_ticks": direcao_ticks,
        "minor_ticks": minor_ticks,
        "cor_moldura": cor_moldura,
        "esp_moldura": esp_moldura,
    }

    # -----------------------------------------------------------------------
    # RENDERIZAÇÃO DO GRÁFICO
    # -----------------------------------------------------------------------
    fig_preview = gerar_figura_matplotlib(
        df=df_editado,
        titulo_grafico=titulo_grafico_input,
        rotulo_x=rotulo_x_input,
        rotulo_y=rotulo_y_input,
        min_y=min_y_val,
        max_y=max_y_val,
        series_ativas=series_ativas,
        cores_periodos=cores_selecionadas,
        estilos_linhas=estilos_linhas,
        fonte_nome=opcao_fonte,
        tamanho_fonte=tamanho_fonte,
        espessura_linha=espessura_linha,
        com_marcadores=com_marcadores,
        bordas_config=bordas_config,
        ticks_config=ticks_config,
        legenda_config=legenda_config,
        exibir_grade=exibir_grade
    )
    st.pyplot(fig_preview)

    # -----------------------------------------------------------------------
    # EXPORTAÇÃO DE ARQUIVOS
    # -----------------------------------------------------------------------
    st.markdown("##### Exportar Resultados:")
    btn_col1, btn_col2, btn_col3 = st.columns(3)

    buf_png = io.BytesIO()
    fig_preview.savefig(buf_png, format="png", dpi=300, bbox_inches="tight")
    btn_col1.download_button(
        label="Download Imagem PNG (300 DPI)",
        data=buf_png.getvalue(),
        file_name=f"{titulo_dia_input.replace(' ', '_').lower()}_figura.png",
        mime="image/png",
        use_container_width=True
    )

    buf_pdf = io.BytesIO()
    fig_preview.savefig(buf_pdf, format="pdf", bbox_inches="tight")
    btn_col2.download_button(
        label="Download Vetor PDF (Publicação)",
        data=buf_pdf.getvalue(),
        file_name=f"{titulo_dia_input.replace(' ', '_').lower()}_figura.pdf",
        mime="application/pdf",
        use_container_width=True
    )

    excel_bytes = gerar_excel_bytes(
        df=df_editado,
        titulo_dia=titulo_dia_input,
        titulo_grafico=titulo_grafico_input,
        min_y=min_y_val,
        max_y=max_y_val
    )
    btn_col3.download_button(
        label="Exportar Planilha Excel (.xlsx)",
        data=excel_bytes,
        file_name=f"{titulo_dia_input.replace(' ', '_').lower()}_dados.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )

    plt.close(fig_preview)

# ---------------------------------------------------------------------------
# BARRA DE SALVAMENTO E GERENCIAMENTO DE PROJETOS (RENDERIZADA NO TOPO)
# ---------------------------------------------------------------------------
with container_barra_projeto:
    with st.container(border=True):
        projetos_salvos_dict = banco_projetos.get("projetos", {})
        nomes_projetos_salvos = list(projetos_salvos_dict.keys())

        p_col1, p_col2, p_col3, p_col4, p_col5 = st.columns([1.8, 1.1, 1.8, 1.1, 0.9])

        with p_col1:
            nome_padrao_proj = st.session_state.get("nome_projeto_atual") or titulo_dia_input
            nome_projeto_input = st.text_input(
                "📌 Nome do Projeto para Salvar:",
                value=nome_padrao_proj,
                placeholder="Ex: Dia 14 de Agosto - Bosque",
                key=f"nome_proj_input_{v}_{nome_padrao_proj}"
            )

        with p_col2:
            st.write("")
            salvar_agora = st.button(
                "💾 Salvar Projeto",
                type="primary",
                use_container_width=True,
                help="Salva todos os dados da tabela e todas as configurações do gráfico para você continuar depois."
            )

        with p_col3:
            opcoes_carregar = ["Selecione um projeto salvo..."] + [
                f"{nome} ({projetos_salvos_dict[nome].get('atualizado_em', '')})"
                for nome in nomes_projetos_salvos
            ]
            projeto_selecionado_label = st.selectbox(
                f"📂 Meus Projetos Salvos ({len(nomes_projetos_salvos)}):",
                opcoes_carregar,
                index=0,
                key=f"sel_proj_salvo_{v}"
            )

        # Extrai o nome real do projeto selecionado no dropdown
        idx_sel_proj = opcoes_carregar.index(projeto_selecionado_label) - 1 if projeto_selecionado_label in opcoes_carregar else -1
        nome_proj_selecionado = nomes_projetos_salvos[idx_sel_proj] if idx_sel_proj >= 0 else None

        with p_col4:
            st.write("")
            payload_atual = construir_payload_projeto(
                nome_projeto=nome_projeto_input or titulo_dia_input,
                df=df_editado,
                titulo_dia=titulo_dia_input,
                titulo_grafico=titulo_grafico_input,
                min_y=min_y_val,
                max_y=max_y_val,
                cfg=st.session_state.cfg
            )
            json_projeto_bytes = json.dumps(payload_atual, ensure_ascii=False, indent=2).encode("utf-8")
            st.download_button(
                label="📥 Baixar (.json)",
                data=json_projeto_bytes,
                file_name=f"{(nome_projeto_input or 'projeto').replace(' ', '_').lower()}_ecoclima.json",
                mime="application/json",
                use_container_width=True,
                help="Baixa um arquivo de backup do projeto no seu computador."
            )

        with p_col5:
            st.write("")
            excluir_proj = st.button(
                "🗑️ Excluir",
                disabled=(nome_proj_selecionado is None and st.session_state.get("nome_projeto_atual") not in projetos_salvos_dict),
                use_container_width=True,
                help="Remove o projeto selecionado da lista de projetos salvos."
            )

        if st.session_state.get("ultimo_salvamento"):
            st.caption(f"✅ Projeto ativo: **{st.session_state.get('nome_projeto_atual')}** — Último salvamento em: `{st.session_state.get('ultimo_salvamento')}`")
        else:
            st.caption("💡 Dica: Clique em **💾 Salvar Projeto** a qualquer momento para guardar seus dados e personalizações e retomar exatamente de onde parou.")

        # Ação: Salvar Projeto
        if salvar_agora:
            nome_final = (nome_projeto_input.strip() if nome_projeto_input and nome_projeto_input.strip() else titulo_dia_input.strip()) or "Meu Projeto"
            novo_payload = construir_payload_projeto(
                nome_projeto=nome_final,
                df=df_editado,
                titulo_dia=titulo_dia_input,
                titulo_grafico=titulo_grafico_input,
                min_y=min_y_val,
                max_y=max_y_val,
                cfg=st.session_state.cfg
            )
            banco_projetos["projetos"][nome_final] = novo_payload
            banco_projetos["ultimo_projeto_ativo"] = nome_final
            gravar_banco_projetos(banco_projetos)

            st.session_state.nome_projeto_atual = nome_final
            st.session_state.ultimo_salvamento = novo_payload["atualizado_em"]
            st.toast(f"Projeto '{nome_final}' salvo com sucesso!", icon="💾")
            st.rerun()

        # Ação: Carregar Projeto Selecionado no Dropdown
        if nome_proj_selecionado is not None:
            proj_para_carregar = projetos_salvos_dict[nome_proj_selecionado]
            aplicar_projeto_na_sessao(proj_para_carregar)
            banco_projetos["ultimo_projeto_ativo"] = nome_proj_selecionado
            gravar_banco_projetos(banco_projetos)
            st.toast(f"Projeto '{nome_proj_selecionado}' carregado!", icon="📂")
            st.rerun()

        # Ação: Excluir Projeto
        if excluir_proj:
            alvo_excluir = nome_proj_selecionado or st.session_state.get("nome_projeto_atual")
            if alvo_excluir and alvo_excluir in banco_projetos.get("projetos", {}):
                del banco_projetos["projetos"][alvo_excluir]
                if banco_projetos.get("ultimo_projeto_ativo") == alvo_excluir:
                    restantes = list(banco_projetos["projetos"].keys())
                    banco_projetos["ultimo_projeto_ativo"] = restantes[-1] if restantes else None
                gravar_banco_projetos(banco_projetos)
                st.session_state.ultimo_salvamento = None
                st.session_state.editor_version += 1
                st.toast(f"Projeto '{alvo_excluir}' removido.", icon="🗑️")
                st.rerun()

# ---------------------------------------------------------------------------
# RESUMO ESTATÍSTICO FORMAL
# ---------------------------------------------------------------------------
st.markdown("---")
with st.expander("Sumário Estatístico Descritivo das Medições", expanded=False):
    s_col1, s_col2, s_col3 = st.columns(3)
    for idx, (periodo, s_col) in enumerate(zip(['Manha', 'Tarde', 'Noite'], [s_col1, s_col2, s_col3])):
        if periodo in df_editado.columns and periodo in series_ativas:
            serie = pd.to_numeric(df_editado[periodo], errors='coerce').dropna()
            if not serie.empty:
                s_col.markdown(f"**Série: {periodo}**")
                s_col.write(f"- Média Amostral: `{serie.mean():.2f} °C`")
                s_col.write(f"- Desvio Padrão: `{serie.std():.2f} °C`")
                s_col.write(f"- Temperatura Mínima: `{serie.min():.2f} °C`")
                s_col.write(f"- Temperatura Máxima: `{serie.max():.2f} °C`")
                s_col.write(f"- Amplitude Térmica: `{(serie.max() - serie.min()):.2f} °C`")
