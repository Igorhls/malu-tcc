"""
=============================================================================
Programa para Geração de Gráficos e Tabelas de Temperatura do Ar - BOSQUE
Projeto: TCC
=============================================================================
Este script reproduz fielmente as tabelas e gráficos dos dias:
- 14 de agosto
- 16 de agosto
- 19 de agosto

Ele gera:
1. Gráficos individuais em alta resolução (PNG 300 DPI e PDF vetorial).
2. Painel combinado com os 3 gráficos para apresentação/artigo/TCC.
3. Planilha Excel (.xlsx) contendo as tabelas formatadas e gráficos nativos.
=============================================================================
"""

import os
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.chart import LineChart, Reference

# ---------------------------------------------------------------------------
# 1. BASE DE DADOS COMPLETA
# ---------------------------------------------------------------------------
PONTOS = [f"P{i}" for i in range(1, 11)]

DADOS = {
    "14": {
        "dia_nome": "Dia 14 de agosto",
        "titulo_grafico": "DIA 14 DE AGOSTO - BOSQUE",
        "ylim": (24, 32),
        "tabela": [
            ("P1", 26.2, 28.5, 25.5),
            ("P2", 26.5, 29.0, 26.1),
            ("P3", 26.5, 29.5, 26.2),
            ("P4", 26.9, 29.6, 26.4),
            ("P5", 27.2, 30.4, 26.9),
            ("P6", 27.4, 29.9, 26.7),
            ("P7", 27.1, 30.9, 26.8),
            ("P8", 27.2, 31.4, 26.9),
            ("P9", 27.3, 30.3, 27.0),
            ("P10", 27.6, 30.5, 27.1),
        ]
    },
    "16": {
        "dia_nome": "Dia 16 de agosto",
        "titulo_grafico": "DIA 16 DE AGOSTO - BOSQUE",
        "ylim": (24, 32),
        "tabela": [
            ("P1", 26.0, 28.5, 25.8),
            ("P2", 25.3, 28.6, 26.0),
            ("P3", 25.3, 28.9, 26.1),
            ("P4", 26.6, 31.0, 26.7),
            ("P5", 26.2, 30.8, 26.8),
            ("P6", 25.5, 29.5, 27.0),
            ("P7", 25.4, 30.4, 27.0),
            ("P8", 26.2, 30.6, 26.9),
            ("P9", 26.0, 30.0, 26.6),
            ("P10", 26.3, 29.9, 27.1),
        ]
    },
    "19": {
        "dia_nome": "Dia 19 de agosto",
        "titulo_grafico": "DIA 19 DE AGOSTO - BOSQUE",
        "ylim": (22, 32),
        "tabela": [
            ("P1", 24.7, 28.2, 27.1),
            ("P2", 24.0, 28.5, 25.5),
            ("P3", 23.4, 28.9, 25.3),
            ("P4", 23.6, 31.5, 26.7),
            ("P5", 23.8, 31.0, 26.2),
            ("P6", 23.6, 28.9, 26.1),
            ("P7", 23.5, 31.0, 25.8),
            ("P8", 23.8, 31.8, 27.0),
            ("P9", 24.2, 30.0, 27.1),
            ("P10", 24.4, 30.6, 25.6),
        ]
    }
}

# ---------------------------------------------------------------------------
# 2. CONFIGURAÇÕES VISUAIS (PALETA FIEL AO EXCEL)
# ---------------------------------------------------------------------------
COR_MANHA = '#265b79'   # Azul petróleo
COR_TARDE = '#e27641'   # Laranja terracota
COR_NOITE = '#236532'   # Verde escuro
COR_TEXTO = '#595959'   # Cinza escuro para títulos e eixos
COR_GRADE = '#e6e6e6'   # Cinza claro para linhas de grade
COR_BORDA = '#d9d9d9'   # Cinza suave para contornos


def configurar_estilo():
    """Configura fontes e propriedades gerais do matplotlib."""
    plt.rcParams['font.sans-serif'] = ['Cambria', 'Times New Roman', 'DejaVu Serif', 'serif']
    plt.rcParams['font.family'] = 'sans-serif'
    plt.rcParams['axes.edgecolor'] = COR_BORDA
    plt.rcParams['axes.linewidth'] = 0.8


def criar_dataframe(chave_dia):
    """Converte a matriz de dados em um DataFrame do pandas."""
    colunas = ['Pontos', 'Manha', 'Tarde', 'Noite']
    return pd.DataFrame(DADOS[chave_dia]['tabela'], columns=colunas)


# ---------------------------------------------------------------------------
# 3. GERAÇÃO DOS GRÁFICOS INDIVIDUAIS
# ---------------------------------------------------------------------------
def gerar_graficos_individuais(pasta_saida="graficos_gerados", com_marcadores=False):
    """Gera um arquivo de imagem (PNG e PDF) para cada dia."""
    os.makedirs(pasta_saida, exist_ok=True)
    configurar_estilo()

    marker_m = 'o' if com_marcadores else None
    marker_t = 's' if com_marcadores else None
    marker_n = '^' if com_marcadores else None

    for dia_key, info in DADOS.items():
        df = criar_dataframe(dia_key)
        fig, ax = plt.subplots(figsize=(7.0, 3.5), dpi=300)

        # Fundo branco e limpo
        fig.patch.set_facecolor('white')
        ax.set_facecolor('white')

        # Linhas de dados
        ax.plot(df['Pontos'], df['Manha'], label='Manha', color=COR_MANHA,
                linewidth=2.0, marker=marker_m, markersize=5)
        ax.plot(df['Pontos'], df['Tarde'], label='Tarde', color=COR_TARDE,
                linewidth=2.0, marker=marker_t, markersize=5)
        ax.plot(df['Pontos'], df['Noite'], label='Noite', color=COR_NOITE,
                linewidth=2.0, marker=marker_n, markersize=5)

        # Título
        ax.set_title(info["titulo_grafico"], fontsize=12, color=COR_TEXTO, pad=26)

        # Legenda no topo centralizada
        leg = ax.legend(loc='upper center', bbox_to_anchor=(0.5, 1.15),
                        ncol=3, frameon=False, fontsize=9.5,
                        handlelength=1.5, handletextpad=0.5, columnspacing=1.6)
        for text in leg.get_texts():
            text.set_color(COR_TEXTO)

        # Rótulos dos eixos
        ax.set_ylabel("Temperatura do ar (°C)", fontsize=9.5, color=COR_TEXTO, labelpad=8)
        ax.set_xlabel("Pontos", fontsize=9.5, color=COR_TEXTO, labelpad=8)

        # Escalas e limites
        ax.set_ylim(info["ylim"])
        ax.yaxis.set_major_locator(ticker.MultipleLocator(2))

        # Grades horizontais sutis
        ax.grid(axis='y', color=COR_GRADE, linestyle='-', linewidth=0.8)
        ax.grid(axis='x', visible=False)

        # Bordas
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        ax.spines['left'].set_visible(False)
        ax.spines['bottom'].set_color(COR_BORDA)
        ax.spines['bottom'].set_linewidth(0.8)

        # Marcadores dos eixos
        ax.tick_params(axis='y', which='both', length=0, labelsize=9.5, labelcolor=COR_TEXTO)
        ax.tick_params(axis='x', which='both', length=3, color=COR_BORDA, labelsize=9.5, labelcolor=COR_TEXTO)

        # Salvar em PNG (alta resolução) e PDF (vetor ideal para Word/LaTeX)
        caminho_png = os.path.join(pasta_saida, f"grafico_dia_{dia_key}.png")
        caminho_pdf = os.path.join(pasta_saida, f"grafico_dia_{dia_key}.pdf")

        plt.savefig(caminho_png, bbox_inches='tight', dpi=300)
        plt.savefig(caminho_pdf, bbox_inches='tight')
        plt.close()

        print(f"[OK] Gráfico gerado: {caminho_png} e {caminho_pdf}")


# ---------------------------------------------------------------------------
# 4. GERAÇÃO DO PAINEL COMBINADO (TODOS OS DIAS)
# ---------------------------------------------------------------------------
def gerar_painel_combinado(pasta_saida="graficos_gerados"):
    """Gera uma figura contendo os 3 gráficos empilhados verticalmente."""
    os.makedirs(pasta_saida, exist_ok=True)
    configurar_estilo()

    fig, axes = plt.subplots(3, 1, figsize=(7.5, 10.5), dpi=300)
    fig.patch.set_facecolor('white')

    for ax, (dia_key, info) in zip(axes, DADOS.items()):
        df = criar_dataframe(dia_key)
        ax.set_facecolor('white')

        ax.plot(df['Pontos'], df['Manha'], label='Manha', color=COR_MANHA, linewidth=2.0)
        ax.plot(df['Pontos'], df['Tarde'], label='Tarde', color=COR_TARDE, linewidth=2.0)
        ax.plot(df['Pontos'], df['Noite'], label='Noite', color=COR_NOITE, linewidth=2.0)

        ax.set_title(info["titulo_grafico"], fontsize=12, color=COR_TEXTO, pad=25)

        leg = ax.legend(loc='upper center', bbox_to_anchor=(0.5, 1.15),
                        ncol=3, frameon=False, fontsize=9.5,
                        handlelength=1.5, handletextpad=0.5, columnspacing=1.6)
        for text in leg.get_texts():
            text.set_color(COR_TEXTO)

        ax.set_ylabel("Temperatura do ar (°C)", fontsize=9.5, color=COR_TEXTO, labelpad=8)
        ax.set_xlabel("Pontos", fontsize=9.5, color=COR_TEXTO, labelpad=8)

        ax.set_ylim(info["ylim"])
        ax.yaxis.set_major_locator(ticker.MultipleLocator(2))

        ax.grid(axis='y', color=COR_GRADE, linestyle='-', linewidth=0.8)
        ax.grid(axis='x', visible=False)

        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        ax.spines['left'].set_visible(False)
        ax.spines['bottom'].set_color(COR_BORDA)
        ax.spines['bottom'].set_linewidth(0.8)

        ax.tick_params(axis='y', which='both', length=0, labelsize=9.5, labelcolor=COR_TEXTO)
        ax.tick_params(axis='x', which='both', length=3, color=COR_BORDA, labelsize=9.5, labelcolor=COR_TEXTO)

    plt.tight_layout(h_pad=3.5)
    caminho_png = os.path.join(pasta_saida, "graficos_todos_dias.png")
    caminho_pdf = os.path.join(pasta_saida, "graficos_todos_dias.pdf")

    plt.savefig(caminho_png, bbox_inches='tight', dpi=300)
    plt.savefig(caminho_pdf, bbox_inches='tight')
    plt.close()

    print(f"[OK] Painel combinado gerado: {caminho_png} e {caminho_pdf}")


# ---------------------------------------------------------------------------
# 5. GERAÇÃO DA PLANILHA EXCEL COM TABELAS E GRÁFICOS NATIVOS
# ---------------------------------------------------------------------------
def gerar_planilha_excel(nome_arquivo="dados_e_graficos_bosque.xlsx"):
    """Cria uma planilha Excel com as tabelas formatadas e gráficos nativos."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOSQUE - Temperaturas"

    borda_fina = Border(
        left=Side(style='thin', color='A6A6A6'),
        right=Side(style='thin', color='A6A6A6'),
        top=Side(style='thin', color='A6A6A6'),
        bottom=Side(style='thin', color='A6A6A6')
    )

    start_row = 2
    for dia_key, bloco in DADOS.items():
        # Cabeçalho do Dia (células mescladas B:E)
        ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row, end_column=5)
        cabecalho = ws.cell(row=start_row, column=2, value=bloco["dia_nome"])
        cabecalho.font = Font(name="Cambria", size=11, bold=True)
        cabecalho.alignment = Alignment(horizontal="center", vertical="center")

        # Cabeçalhos das colunas
        colunas = ["Pontos", "Manha", "Tarde", "Noite"]
        for col_idx, col_name in enumerate(colunas, start=2):
            c = ws.cell(row=start_row + 1, column=col_idx, value=col_name)
            c.font = Font(name="Cambria", size=10, bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = borda_fina

        for c_idx in range(2, 6):
            ws.cell(row=start_row, column=c_idx).border = borda_fina

        # Preenchimento das linhas da tabela
        inicio_dados = start_row + 2
        for r_idx, linha in enumerate(bloco["tabela"]):
            linha_atual = inicio_dados + r_idx
            for col_idx, valor in enumerate(linha, start=2):
                c = ws.cell(row=linha_atual, column=col_idx, value=valor)
                c.font = Font(name="Cambria", size=10)
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.border = borda_fina
                if isinstance(valor, float):
                    c.number_format = '0.0'

        fim_dados = inicio_dados + len(bloco["tabela"]) - 1

        # Criação do gráfico nativo do Excel
        chart = LineChart()
        chart.title = bloco["titulo_grafico"]
        chart.style = 10
        chart.y_axis.title = "Temperatura do ar (°C)"
        chart.x_axis.title = "Pontos"
        chart.y_axis.scaling.min = bloco["ylim"][0]
        chart.y_axis.scaling.max = bloco["ylim"][1]
        chart.y_axis.majorUnit = 2
        chart.width = 16
        chart.height = 9.5

        dados_ref = Reference(ws, min_col=3, min_row=start_row + 1, max_col=5, max_row=fim_dados)
        categorias_ref = Reference(ws, min_col=2, min_row=inicio_dados, max_row=fim_dados)
        chart.add_data(dados_ref, titles_from_data=True)
        chart.set_categories(categorias_ref)

        ws.add_chart(chart, f"G{start_row}")

        start_row = fim_dados + 3

    # Ajuste da largura das colunas para visualização agradável
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 10
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 10
    ws.column_dimensions['F'].width = 4

    wb.save(nome_arquivo)
    print(f"[OK] Planilha Excel gerada: {nome_arquivo}")


# ---------------------------------------------------------------------------
# 6. FUNÇÃO PRINCIPAL DE EXECUÇÃO
# ---------------------------------------------------------------------------
def main():
    print("=" * 60)
    print("GERADOR DE GRÁFICOS E TABELAS - BOSQUE (TCC)")
    print("=" * 60)

    # 1. Gráficos individuais
    print("\n1. Gerando gráficos individuais (PNG 300 DPI e PDF)...")
    gerar_graficos_individuais()

    # 2. Painel combinado
    print("\n2. Gerando painel combinado com todos os gráficos...")
    gerar_painel_combinado()

    # 3. Planilha Excel
    print("\n3. Gerando planilha Excel com tabelas e gráficos embutidos...")
    gerar_planilha_excel()

    print("\n" + "=" * 60)
    print("Concluído com sucesso! Verifique a pasta 'graficos_gerados' e o arquivo Excel.")
    print("=" * 60)


if __name__ == "__main__":
    main()
